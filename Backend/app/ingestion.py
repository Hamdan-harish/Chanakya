"""Format-aware, provenance-preserving document ingestion."""

from __future__ import annotations

import csv
import json
from enum import StrEnum
from pathlib import Path
from typing import Any

import fitz
from docx import Document
from openpyxl import load_workbook
from pydantic import BaseModel, Field


class InputFormat(StrEnum):
    PDF = "pdf"
    DOCX = "docx"
    TXT = "txt"
    CSV = "csv"
    XLSX = "xlsx"
    JSON = "json"


class SourceLocation(BaseModel):
    """Original location of a normalized fragment."""

    page_number: int | None = Field(default=None, ge=1)
    paragraph_number: int | None = Field(default=None, ge=1)
    table_number: int | None = Field(default=None, ge=1)
    line_number: int | None = Field(default=None, ge=1)
    sheet_name: str | None = None
    row_number: int | None = Field(default=None, ge=1)
    column_number: int | None = Field(default=None, ge=1)
    json_path: str | None = None
    character_start: int | None = Field(default=None, ge=0)
    character_end: int | None = Field(default=None, ge=0)


class NormalizedFragment(BaseModel):
    """Source content plus a stable location for future extraction stages."""

    content: str
    original_value: Any
    location: SourceLocation


class NormalizedDocument(BaseModel):
    """Common, loss-conscious representation of a single input file."""

    document_id: str
    format: InputFormat
    fragments: list[NormalizedFragment] = Field(default_factory=list)
    requires_ocr: bool = False
    pages_requiring_ocr: list[int] = Field(default_factory=list)


class UnsupportedFormatError(ValueError):
    """Raised when a file suffix has no ingestion parser."""


_FORMAT_BY_SUFFIX = {
    ".pdf": InputFormat.PDF,
    ".docx": InputFormat.DOCX,
    ".txt": InputFormat.TXT,
    ".csv": InputFormat.CSV,
    ".xlsx": InputFormat.XLSX,
    ".json": InputFormat.JSON,
}


def detect_format(path: Path) -> InputFormat:
    """Detect a supported input format from its file extension."""

    try:
        return _FORMAT_BY_SUFFIX[path.suffix.lower()]
    except KeyError as error:
        raise UnsupportedFormatError(f"Unsupported input format: {path.suffix or '<none>'}") from error


def ingest_file(path: str | Path, document_id: str | None = None) -> NormalizedDocument:
    """Parse one supported file into normalized, provenance-bearing fragments."""

    source_path = Path(path)
    input_format = detect_format(source_path)
    normalized_id = document_id or source_path.stem

    if input_format is InputFormat.PDF:
        return _ingest_pdf(source_path, normalized_id)
    if input_format is InputFormat.DOCX:
        return _ingest_docx(source_path, normalized_id)
    if input_format is InputFormat.TXT:
        return _ingest_txt(source_path, normalized_id)
    if input_format is InputFormat.CSV:
        return _ingest_csv(source_path, normalized_id)
    if input_format is InputFormat.XLSX:
        return _ingest_xlsx(source_path, normalized_id)
    return _ingest_json(source_path, normalized_id)


def _ingest_pdf(path: Path, document_id: str) -> NormalizedDocument:
    fragments: list[NormalizedFragment] = []
    pages_requiring_ocr: list[int] = []

    with fitz.open(path) as pdf:
        for page_number, page in enumerate(pdf, start=1):
            text = page.get_text("text")
            if not text.strip():
                pages_requiring_ocr.append(page_number)
                continue
            fragments.append(
                NormalizedFragment(
                    content=text,
                    original_value=text,
                    location=SourceLocation(
                        page_number=page_number,
                        character_start=0,
                        character_end=len(text),
                    ),
                )
            )

    return NormalizedDocument(
        document_id=document_id,
        format=InputFormat.PDF,
        fragments=fragments,
        requires_ocr=bool(pages_requiring_ocr),
        pages_requiring_ocr=pages_requiring_ocr,
    )


def _ingest_docx(path: Path, document_id: str) -> NormalizedDocument:
    document = Document(path)
    fragments: list[NormalizedFragment] = []

    for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
        if paragraph.text:
            fragments.append(
                NormalizedFragment(
                    content=paragraph.text,
                    original_value=paragraph.text,
                    location=SourceLocation(paragraph_number=paragraph_number),
                )
            )

    for table_number, table in enumerate(document.tables, start=1):
        for row_number, row in enumerate(table.rows, start=1):
            for column_number, cell in enumerate(row.cells, start=1):
                if cell.text:
                    fragments.append(
                        NormalizedFragment(
                            content=cell.text,
                            original_value=cell.text,
                            location=SourceLocation(
                                table_number=table_number,
                                row_number=row_number,
                                column_number=column_number,
                            ),
                        )
                    )

    return NormalizedDocument(document_id=document_id, format=InputFormat.DOCX, fragments=fragments)


def _ingest_txt(path: Path, document_id: str) -> NormalizedDocument:
    source = path.read_text(encoding="utf-8")
    fragments: list[NormalizedFragment] = []
    offset = 0

    for line_number, raw_line in enumerate(source.splitlines(keepends=True), start=1):
        content = raw_line.rstrip("\r\n")
        fragments.append(
            NormalizedFragment(
                content=content,
                original_value=content,
                location=SourceLocation(
                    line_number=line_number,
                    character_start=offset,
                    character_end=offset + len(content),
                ),
            )
        )
        offset += len(raw_line)

    return NormalizedDocument(document_id=document_id, format=InputFormat.TXT, fragments=fragments)


def _ingest_csv(path: Path, document_id: str) -> NormalizedDocument:
    fragments: list[NormalizedFragment] = []
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        for row_number, row in enumerate(csv.reader(source), start=1):
            for column_number, value in enumerate(row, start=1):
                fragments.append(
                    NormalizedFragment(
                        content=value,
                        original_value=value,
                        location=SourceLocation(row_number=row_number, column_number=column_number),
                    )
                )
    return NormalizedDocument(document_id=document_id, format=InputFormat.CSV, fragments=fragments)


def _ingest_xlsx(path: Path, document_id: str) -> NormalizedDocument:
    workbook = load_workbook(path, read_only=True, data_only=False)
    fragments: list[NormalizedFragment] = []
    try:
        for worksheet in workbook.worksheets:
            for row_number, row in enumerate(worksheet.iter_rows(), start=1):
                for column_number, cell in enumerate(row, start=1):
                    if cell.value is not None:
                        fragments.append(
                            NormalizedFragment(
                                content=str(cell.value),
                                original_value=cell.value,
                                location=SourceLocation(
                                    sheet_name=worksheet.title,
                                    row_number=row_number,
                                    column_number=column_number,
                                ),
                            )
                        )
    finally:
        workbook.close()
    return NormalizedDocument(document_id=document_id, format=InputFormat.XLSX, fragments=fragments)


def _ingest_json(path: Path, document_id: str) -> NormalizedDocument:
    value = json.loads(path.read_text(encoding="utf-8"))
    fragments: list[NormalizedFragment] = []
    for json_path, leaf in _json_leaves(value):
        fragments.append(
            NormalizedFragment(
                content=json.dumps(leaf, ensure_ascii=False) if not isinstance(leaf, str) else leaf,
                original_value=leaf,
                location=SourceLocation(json_path=json_path),
            )
        )
    return NormalizedDocument(document_id=document_id, format=InputFormat.JSON, fragments=fragments)


def _json_leaves(value: Any, path: str = "$") -> list[tuple[str, Any]]:
    if isinstance(value, dict):
        leaves: list[tuple[str, Any]] = []
        for key, child in value.items():
            escaped_key = str(key).replace("\\", "\\\\").replace('"', '\\"')
            leaves.extend(_json_leaves(child, f'{path}["{escaped_key}"]'))
        return leaves
    if isinstance(value, list):
        leaves = []
        for index, child in enumerate(value):
            leaves.extend(_json_leaves(child, f"{path}[{index}]"))
        return leaves
    return [(path, value)]
