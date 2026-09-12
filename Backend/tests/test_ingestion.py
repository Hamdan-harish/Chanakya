import json

import pymupdf as fitz
from docx import Document
from openpyxl import Workbook

from app.ingestion import InputFormat, ingest_file


def test_pdf_preserves_page_provenance_and_marks_empty_pages(tmp_path) -> None:
    path = tmp_path / "report.pdf"
    pdf = fitz.open()
    first_page = pdf.new_page()
    first_page.insert_text((72, 72), "Native report text")
    pdf.new_page()
    pdf.save(path)
    pdf.close()

    result = ingest_file(path, document_id="pdf-1")

    assert result.format is InputFormat.PDF
    assert result.fragments[0].location.page_number == 1
    assert result.fragments[0].content == "Native report text\n"
    assert result.requires_ocr is True
    assert result.pages_requiring_ocr == [2]


def test_docx_preserves_paragraph_and_table_locations(tmp_path) -> None:
    path = tmp_path / "report.docx"
    document = Document()
    document.add_paragraph("Witness statement")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Field"
    table.cell(0, 1).text = "Value"
    document.save(path)

    result = ingest_file(path)

    assert result.format is InputFormat.DOCX
    assert result.fragments[0].location.paragraph_number == 1
    table_fragment = next(item for item in result.fragments if item.content == "Value")
    assert (table_fragment.location.table_number, table_fragment.location.row_number) == (1, 1)
    assert table_fragment.location.column_number == 2


def test_txt_preserves_lines_and_offsets(tmp_path) -> None:
    path = tmp_path / "report.txt"
    path.write_text("first\nsecond", encoding="utf-8")

    result = ingest_file(path)

    assert result.format is InputFormat.TXT
    assert result.fragments[1].content == "second"
    assert result.fragments[1].location.line_number == 2
    assert result.fragments[1].location.character_start == 6


def test_csv_preserves_each_value_and_cell_location(tmp_path) -> None:
    path = tmp_path / "cdr.csv"
    path.write_text("caller,callee\n111,222\n", encoding="utf-8")

    result = ingest_file(path)

    value = next(item for item in result.fragments if item.content == "222")
    assert result.format is InputFormat.CSV
    assert value.original_value == "222"
    assert (value.location.row_number, value.location.column_number) == (2, 2)


def test_xlsx_preserves_sheet_formula_and_cell_location(tmp_path) -> None:
    path = tmp_path / "transactions.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Transfers"
    worksheet["A1"] = "Amount"
    worksheet["B2"] = "=1+1"
    workbook.save(path)

    result = ingest_file(path)

    formula = next(item for item in result.fragments if item.content == "=1+1")
    assert result.format is InputFormat.XLSX
    assert formula.original_value == "=1+1"
    assert (formula.location.sheet_name, formula.location.row_number, formula.location.column_number) == (
        "Transfers",
        2,
        2,
    )


def test_json_preserves_paths_and_original_values(tmp_path) -> None:
    path = tmp_path / "intel.json"
    path.write_text(json.dumps({"profiles": [{"handle": "@source", "active": True}]}), encoding="utf-8")

    result = ingest_file(path)

    handle = next(item for item in result.fragments if item.content == "@source")
    active = next(item for item in result.fragments if item.original_value is True)
    assert result.format is InputFormat.JSON
    assert handle.location.json_path == '$["profiles"][0]["handle"]'
    assert active.location.json_path == '$["profiles"][0]["active"]'
