import json
from io import BytesIO
from typing import Any

import pymupdf as fitz
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

import app.main as main_module
from app.ingestion import NormalizedDocument
from app.ocr import BoundingBox, OCRFragment, OCRPageResult, Point
from app.processing import EvidenceProcessingPipeline
from app.relationship_extraction import RelationExtractionResult
from app.relationship_extraction import DeBERTaRelationClassifier
from app.schemas import Entity, Provenance, Relationship


class FakeEntityExtractor:
    def __init__(self) -> None:
        self.documents: list[NormalizedDocument] = []

    def extract(self, document: NormalizedDocument) -> list[Entity]:
        self.documents.append(document)
        entities: list[Entity] = []
        for fragment_index, fragment in enumerate(document.fragments):
            for surface in ("Rahul", "Ahmed"):
                start = fragment.content.find(surface)
                if start >= 0:
                    entities.append(
                        Entity(
                            id=f"{document.document_id}:entity:{len(entities) + 1}",
                            type="person",
                            label=surface,
                            confidence=0.9,
                            provenance=[
                                Provenance(
                                    document_id=document.document_id,
                                    page_number=fragment.location.page_number,
                                    source_location=json.dumps(
                                        fragment.location.model_dump(exclude_none=True)
                                    ),
                                    evidence_text=fragment.content,
                                )
                            ],
                            attributes={
                                "fragment_index": fragment_index,
                                "start": start,
                                "end": start + len(surface),
                            },
                        )
                    )
        return entities


class FakeRelationshipExtractor:
    def __init__(self, relation_type: str = "CONTACTED") -> None:
        self.calls = 0
        self.relation_type = relation_type

    def extract_with_decisions(
        self, document: NormalizedDocument, entities: list[Entity]
    ) -> RelationExtractionResult:
        self.calls += 1
        if len(entities) < 2:
            return RelationExtractionResult()
        fragment_index = entities[0].attributes["fragment_index"]
        fragment = document.fragments[fragment_index]
        relationship = Relationship(
            id=f"{document.document_id}:relationship:1",
            source_entity_id=entities[0].id,
            target_entity_id=entities[1].id,
            type=self.relation_type,
            confidence=0.9,
            provenance=[
                Provenance(
                    document_id=document.document_id,
                    page_number=fragment.location.page_number,
                    source_location=json.dumps(fragment.location.model_dump(exclude_none=True)),
                    evidence_text=fragment.content,
                )
            ],
        )
        return RelationExtractionResult(relationships=[relationship])


class FakeOCR:
    def __init__(self) -> None:
        self.image_calls = 0
        self.pdf_calls = 0

    def _page(self, document_id: str, page_number: int | None, source: str) -> OCRPageResult:
        return OCRPageResult(
            document_id=document_id,
            page_number=page_number,
            source_path=source,
            image_width=640,
            image_height=480,
            fragments=[
                OCRFragment(
                    document_id=document_id,
                    page_number=page_number,
                    text="Rahul contacted Ahmed.",
                    evidence_text="Rahul contacted Ahmed.",
                    confidence=0.95,
                    bounding_box=BoundingBox(
                        points=[
                            Point(x=1, y=1),
                            Point(x=200, y=1),
                            Point(x=200, y=30),
                            Point(x=1, y=30),
                        ]
                    ),
                    source_path=source,
                    image_width=640,
                    image_height=480,
                )
            ],
        )

    def extract_image(self, path: Any, *, document_id: str) -> OCRPageResult:
        self.image_calls += 1
        return self._page(document_id, None, str(path))

    def extract_required_pdf_pages(
        self, path: Any, document: NormalizedDocument
    ) -> list[OCRPageResult]:
        self.pdf_calls += 1
        return [self._page(document.document_id, page, str(path)) for page in document.pages_requiring_ocr]


def _install_pipeline(monkeypatch, *, relation_type: str = "CONTACTED"):
    entity = FakeEntityExtractor()
    relation = FakeRelationshipExtractor(relation_type)
    ocr = FakeOCR()
    pipeline = EvidenceProcessingPipeline(
        entity_extractor=entity,
        relationship_extractor=relation,
        ocr_service=ocr,
    )
    monkeypatch.setattr(main_module, "processing_pipeline", pipeline)
    return TestClient(main_module.app), entity, relation, ocr


def _post(client: TestClient, filename: str, data: bytes, content_type: str):
    return client.post("/process", files={"artifact": (filename, data, content_type)})


def _pdf_bytes(*, text: str | None) -> bytes:
    pdf = fitz.open()
    page = pdf.new_page()
    if text:
        page.insert_text((72, 72), text)
    data = pdf.tobytes()
    pdf.close()
    return data


def test_txt_runs_narrative_entity_and_relationship_pipeline(monkeypatch) -> None:
    client, entity, relation, _ = _install_pipeline(monkeypatch)
    response = _post(client, "report.txt", b"Rahul contacted Ahmed.", "text/plain")

    assert response.status_code == 200
    body = response.json()
    assert body["processing"] == {
        "input_format": "txt",
        "mode": "narrative",
        "ocr_used": False,
        "docres_used": False,
    }
    assert len(body["fragments"]) == 1
    assert len(body["entities"]) == 2
    assert [item["type"] for item in body["relationships"]] == ["CONTACTED"]
    assert entity.documents and relation.calls == 1


def test_docx_uses_native_ingestion_and_paragraph_provenance(monkeypatch) -> None:
    client, _, _, ocr = _install_pipeline(monkeypatch)
    output = BytesIO()
    document = Document()
    document.add_paragraph("Rahul contacted Ahmed.")
    document.save(output)

    response = _post(
        client,
        "statement.docx",
        output.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert response.status_code == 200
    assert response.json()["fragments"][0]["location"]["paragraph_number"] == 1
    assert ocr.image_calls == ocr.pdf_calls == 0


def test_native_pdf_does_not_call_ocr(monkeypatch) -> None:
    client, _, _, ocr = _install_pipeline(monkeypatch)
    response = _post(client, "report.pdf", _pdf_bytes(text="Rahul contacted Ahmed."), "application/pdf")
    assert response.status_code == 200
    assert response.json()["processing"]["ocr_used"] is False
    assert ocr.pdf_calls == 0


def test_image_uses_ocr_and_preserves_pixel_provenance(monkeypatch) -> None:
    client, _, _, ocr = _install_pipeline(monkeypatch)
    response = _post(client, "scan.png", b"synthetic-image", "image/png")
    body = response.json()
    assert response.status_code == 200
    assert body["processing"]["ocr_used"] is True
    assert body["fragments"][0]["location"]["image_width"] == 640
    assert len(body["fragments"][0]["location"]["bounding_box"]) == 4
    assert body["fragments"][0]["location"]["source_path"] == "scan.png"
    assert ocr.image_calls == 1


def test_scanned_pdf_uses_ocr_only_for_required_page(monkeypatch) -> None:
    client, _, _, ocr = _install_pipeline(monkeypatch)
    response = _post(client, "scan.pdf", _pdf_bytes(text=None), "application/pdf")
    body = response.json()
    assert response.status_code == 200
    assert body["processing"]["ocr_used"] is True
    assert body["fragments"][0]["location"]["page_number"] == 1
    assert ocr.pdf_calls == 1


def test_csv_cdr_creates_deterministic_contacted_with_cell_provenance(monkeypatch) -> None:
    client, entity, relation, _ = _install_pipeline(monkeypatch)
    response = _post(client, "cdr.csv", b"caller,receiver,date\n111,222,2026-08-12\n", "text/csv")
    body = response.json()
    edge = body["relationships"][0]
    assert edge["type"] == "CONTACTED"
    assert edge["attributes"]["source_location"]["row_number"] == 2
    assert edge["attributes"]["source_location"]["column_number"] == 1
    assert edge["attributes"]["target_location"]["column_number"] == 2
    assert entity.documents == [] and relation.calls == 0


def test_csv_transfer_creates_deterministic_transfer(monkeypatch) -> None:
    client, _, _, _ = _install_pipeline(monkeypatch)
    response = _post(
        client,
        "transfer.csv",
        b"sender_account,receiver_account,amount\nACCT001,ACCT002,50000\n",
        "text/csv",
    )
    body = response.json()
    assert [item["type"] for item in body["relationships"]] == ["TRANSFERRED_TO"]
    assert [item["type"] for item in body["entities"]] == ["account_number", "account_number"]


def test_xlsx_preserves_sheet_and_creates_known_relation(monkeypatch) -> None:
    client, _, _, _ = _install_pipeline(monkeypatch)
    output = BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Calls"
    sheet.append(["caller", "receiver"])
    sheet.append(["111", "222"])
    workbook.save(output)
    response = _post(
        client,
        "calls.xlsx",
        output.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    body = response.json()
    assert body["relationships"][0]["type"] == "CONTACTED"
    assert body["relationships"][0]["attributes"]["source_location"]["sheet_name"] == "Calls"


def test_json_known_schema_preserves_paths(monkeypatch) -> None:
    client, _, _, _ = _install_pipeline(monkeypatch)
    data = json.dumps({"calls": [{"caller": "111", "receiver": "222"}]}).encode()
    response = _post(client, "calls.json", data, "application/json")
    body = response.json()
    assert body["relationships"][0]["type"] == "CONTACTED"
    assert body["relationships"][0]["attributes"]["source_location"]["json_path"] == '$["calls"][0]["caller"]'


def test_unknown_structured_schema_preserves_fragments_without_edges(monkeypatch) -> None:
    client, _, _, _ = _install_pipeline(monkeypatch)
    response = _post(client, "unknown.csv", b"name,note\nRahul,witness\n", "text/csv")
    body = response.json()
    assert len(body["fragments"]) == 4
    assert body["entities"] == []
    assert body["relationships"] == []


def test_unsupported_extension_is_rejected(monkeypatch) -> None:
    client, _, _, _ = _install_pipeline(monkeypatch)
    response = _post(client, "archive.zip", b"not-a-zip", "application/zip")
    assert response.status_code == 415


def test_content_type_must_match_extension(monkeypatch) -> None:
    client, _, _, _ = _install_pipeline(monkeypatch)
    response = _post(client, "report.txt", b"text", "image/png")
    assert response.status_code == 415


def test_empty_and_malformed_files_are_rejected(monkeypatch) -> None:
    client, _, _, _ = _install_pipeline(monkeypatch)
    assert _post(client, "empty.txt", b"", "text/plain").status_code == 400
    assert _post(client, "broken.json", b"{", "application/json").status_code == 422


def test_no_relation_is_never_returned_as_graph_edge(monkeypatch) -> None:
    client, _, _, _ = _install_pipeline(monkeypatch, relation_type="NO_RELATION")
    response = _post(client, "report.txt", b"Rahul and Ahmed were listed.", "text/plain")
    assert response.status_code == 200
    assert response.json()["relationships"] == []


def test_accepted_relationship_is_traceable_to_exact_evidence(monkeypatch) -> None:
    client, _, _, _ = _install_pipeline(monkeypatch)
    response = _post(client, "report.txt", b"Rahul contacted Ahmed.", "text/plain")
    edge = response.json()["relationships"][0]
    assert edge["provenance"][0]["evidence_text"] == "Rahul contacted Ahmed."
    location = json.loads(edge["provenance"][0]["source_location"])
    assert location["line_number"] == 1
    assert location["source_path"] == "report.txt"


def test_relation_model_path_is_configurable_without_loading_model(monkeypatch, tmp_path) -> None:
    monkeypatch.setenv("CHANAKYA_RELATION_MODEL_DIR", str(tmp_path))
    classifier = DeBERTaRelationClassifier()
    assert classifier.model_path == tmp_path.resolve()
    assert classifier._model is None
